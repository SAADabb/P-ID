"""One-sheet drawing, DEXPI graphics, decision register, line numbers and engineering lists."""

from __future__ import annotations

import copy
from pathlib import Path

import pytest
import yaml
from lxml import etree
from openpyxl import load_workbook
from pypdf import PdfReader

from pid_workflow import review
from pid_workflow.decisions import build_register, load_plan, sheet_content
from pid_workflow.enrich import enrich
from pid_workflow.export_dexpi import export_topology
from pid_workflow.lists import write_xlsx
from pid_workflow.onesheet import build_onesheet, onesheet_qa
from pid_workflow.proteus_graphics import add_graphics, verify_graphics
from pid_workflow.render_onesheet import draw_onesheet, save_onesheet_pdf
from pid_workflow.render_pdf import Style
from pid_workflow.roundtrip import verify_roundtrip
from pid_workflow.schema import logical

from .conftest import PROJECT, ROOT


@pytest.fixture(scope="module")
def sheet(enriched, hints):
    topo = enriched[0]
    reg = build_register(topo, load_plan(PROJECT / "review_plan.yaml"))
    os_ = build_onesheet(topo, hints)
    d = draw_onesheet(topo, os_, Style(ROOT / "policies" / "rendering.yaml"), sheet_content(topo, reg))
    return topo, os_, d, reg


def test_one_sheet_layout_is_clean_and_connected(sheet):
    topo, os_, d, _ = sheet
    assert onesheet_qa(os_) == []
    assert len(os_.zones) == 8
    # every connection between areas is drawn as one routed pipe (no off-page connectors)
    assert os_.links and all(lk.edge.id in os_.routes for lk in os_.links)
    assert not any(n.kind == "offpage" for n in os_.layout.nodes.values())
    # every routed pipe is orthogonal
    for path in os_.routes.values():
        for (x1, y1), (x2, y2) in zip(path, path[1:]):
            assert abs(x1 - x2) < 1e-6 or abs(y1 - y2) < 1e-6
    assert 0.75 <= d.scale <= 1.2


def test_every_object_is_drawn_once_on_the_sheet(sheet):
    topo, os_, d, _ = sheet
    symbols = {(p.owner.kind, p.owner.ref) for p in d.prims if p.owner.role == "symbol"}
    assert {e.id for e in topo.equipment} <= {r for k, r in symbols if k == "equipment"}
    assert {c.id for c in topo.components if c.kind != "manifold"} <= {r for k, r in symbols if k == "valve"}
    assert {i.id for i in topo.instruments} <= {r for k, r in symbols if k == "instrument"}


def test_pdf_is_a_single_a0_page(sheet, tmp_path: Path):
    topo, _, d, _ = sheet
    out = tmp_path / "sheet.pdf"
    save_onesheet_pdf(d, topo, out)
    pages = PdfReader(str(out)).pages
    assert len(pages) == 1
    w, h = float(pages[0].mediabox.width), float(pages[0].mediabox.height)
    assert abs(w / 72 * 25.4 - 1189) < 1 and abs(h / 72 * 25.4 - 841) < 1
    text = pages[0].extract_text()
    for must in ("P-030", "TK-001", "LEGEND", "HOLDS", "PID-LPF-BSST-001", "CW-1"):
        assert must in text, must


def test_dexpi_file_carries_the_drawing(sheet, tmp_path: Path):
    topo, os_, d, _ = sheet
    xml = tmp_path / "g.proteus.xml"
    export_topology(topo, xml)
    stats = add_graphics(xml, d, os_, topo, None)
    assert stats["placed"] == len(topo.equipment) + len([c for c in topo.components if c.kind != "manifold"]) + len(topo.instruments) + len(topo.boundaries)
    assert stats["centerlines"] == len(topo.process_lines)
    checks = verify_graphics(xml, stats)
    assert all(ok for _, ok, _ in checks), [c for c in checks if not c[1]]
    # the model part is unchanged by the graphics
    assert verify_roundtrip(topo, xml).ok
    root = etree.parse(str(xml)).getroot()
    names = {s.get("ComponentName") for s in root.find("ShapeCatalogue")}
    used = [e.get("ComponentName") for e in root.iter("Equipment", "PipingComponent", "ProcessInstrumentationFunction", "PipeOffPageConnector")
            if e.get("ComponentName")]
    assert used and set(used) <= names
    # pipes start/end where they are drawn: every centre line has at least two coordinates
    assert all(len(cl.findall("Coordinate")) >= 2 for cl in root.iter("CenterLine") if cl.find("Presentation") is not None)


def test_line_numbers_per_logical_line_and_idempotent(enriched, policy):
    topo = enriched[0]
    by_logical: dict[str, set[str | None]] = {}
    for ln in topo.process_lines:
        by_logical.setdefault(logical(ln), set()).add(ln.line_number)
    assert all(len(v) == 1 for v in by_logical.values())
    numbered = [ln for ln in topo.process_lines if ln.line_number]
    assert numbered and all(ln.line_number_status == "proposed" for ln in numbered)
    assert not any(ln.line_number for ln in topo.process_lines if ln.service.startswith(("Drain", "Vent")))
    # numbers are unique per logical line
    first = {lg: next(iter(v)) for lg, v in by_logical.items() if next(iter(v))}
    assert len(set(first.values())) == len(first)
    twin = copy.deepcopy(topo)
    enrich(twin, policy)
    assert [ln.line_number for ln in twin.process_lines] == [ln.line_number for ln in topo.process_lines]


def test_decision_register_reduces_review_to_holds_and_rules(enriched):
    topo = enriched[0]
    plan = load_plan(PROJECT / "review_plan.yaml")
    reg = build_register(topo, plan)
    open_urs = [u for u in topo.unresolved_requirements if u.status == "open"]
    # nothing is dropped: every open point is either a hold or datasheet data
    assert len(reg.holds) + len(reg.datasheet) == len(open_urs)
    assert [h.hold.num for h in reg.holds] == list(range(1, len(reg.holds) + 1))
    assert all(h.drawn_as for h in reg.holds)
    assert {r for r, _, _ in reg.rules} >= {"pump.auxiliaries", "lines.numbering"}
    # answering a hold closes it
    plan2 = copy.deepcopy(plan)
    plan2["holds"][0]["answer"] = "accepted"
    reg2 = build_register(topo, plan2)
    assert len(reg2.holds) == len(reg.holds) - 1 and len(reg2.closed) == 1
    # without a plan every non-missing-value open point is a hold
    reg3 = build_register(topo, {})
    assert all(h.category != "missing_value" for h in reg3.holds)


def test_rule_approval_includes_its_changes_to_source_objects(enriched):
    topo = copy.deepcopy(enriched[0])
    res = review.approve(topo, rule="lines.numbering")
    assert res.modifications > 0
    assert not [m for m in topo.modifications if m.provenance.rule == "lines.numbering" and m.provenance.status == "pending_review"]
    assert "lines.numbering" not in {r for r, _, _ in build_register(topo, {}).rules}


def test_engineering_lists_workbook(enriched, tmp_path: Path):
    topo = enriched[0]
    reg = build_register(topo, load_plan(PROJECT / "review_plan.yaml"))
    wb = load_workbook(write_xlsx(topo, reg, tmp_path / "lists.xlsx"))
    assert wb.sheetnames == ["Equipment list", "Line list", "Valve list", "Instrument index", "Holds", "Rule approvals", "Datasheet data"]
    assert wb["Equipment list"].max_row == len(topo.equipment) + 1
    assert wb["Valve list"].max_row == len([c for c in topo.components if c.kind != "manifold"]) + 1
    assert wb["Instrument index"].max_row == len(topo.instruments) + 1
    tags = {r[0] for r in wb["Equipment list"].iter_rows(min_row=2, values_only=True)}
    assert {"P-030", "TK-001", "ECH-030"} <= tags


def test_onesheet_hints_are_optional(enriched, hints):
    """Without the explicit row arrangement the zones are arranged by the search."""
    h = {k: v for k, v in yaml.safe_load(yaml.safe_dump(hints)).items() if k != "onesheet"}
    os_ = build_onesheet(enriched[0], h)
    assert [p for p in onesheet_qa(os_) if "shares a track" not in p] == []  # searched arrangement: no overlaps
