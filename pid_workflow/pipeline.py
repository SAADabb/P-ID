"""Stage orchestration: enrichment -> validation -> export -> round trip -> layout -> PDF -> reports."""

from __future__ import annotations

import copy
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import yaml

from . import reports as R
from .enrich import EnrichResult, enrich
from .export_dexpi import ExportResult, export_topology
from .decisions import Register, build_register, load_plan, register_md, sheet_content
from .layout import DrawingLayout, build_layouts, layout_qa
from .lists import write_xlsx
from .onesheet import OneSheet, build_onesheet, onesheet_qa
from .proteus_graphics import add_graphics, render_svg, verify_graphics
from .render_onesheet import draw_onesheet, save_onesheet_pdf
from .policy import Policy, load_policy
from .render_pdf import Style, render_review_pdf
from .roundtrip import RoundTripReport, render_markdown as roundtrip_md, verify_roundtrip
from .schema import Topology
from .topology_io import dumps_topology, load_topology, save_topology, sha256_file
from .validate import ValidationReport, render_markdown as validation_md, validate


@dataclass
class PFDResult:
    pfd: Any
    qa: list[str]
    scale: float
    roundtrip_ok: bool
    files: list[Path]


def run_pfd(project: str | Path) -> PFDResult:
    """PFD from the source topology (no P&ID rules): A0 sheet, stream table, equipment summary, DEXPI with graphics."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    from .pfd import STREAM_ROWS, build_pfd, stream_labels

    proj = load_project(project)
    source = load_topology(proj.p("source_topology"))
    pfd = build_pfd(source)
    t = pfd.topo
    reg = build_register(t, load_plan(proj.plan_path))
    content = sheet_content(t, reg)
    content.kind = "pfd"
    content.line_labels = stream_labels(pfd)
    content.tbd = []
    content.equipment_summary = pfd.equipment_summary
    content.notes = [
        "Process flow diagram derived from the same extraction as the P&ID: major equipment and process streams only; "
        "isolation, check, drain and vent valves, strainers and local instruments are on the P&ID.",
        "Identical parallel units are drawn once; their tags and number are given on the symbol.",
        "Stream data are those stated in the specification; '*' = carried over from equipment data (e.g. pump design flow). "
        "Blank = not given, not assumed.",
    ] + [n for n in content.notes if "Orange" not in n and "Line numbers" not in n]
    rows = [["Stream"] + [str(s_.no) for s_ in pfd.streams]]
    for r in STREAM_ROWS:
        rows.append([r] + [s_.values.get(r, "") + ("*" if r in s_.derived else "") for s_ in pfd.streams])
    content.stream_table = rows
    hints = _hints(proj)
    os_ = build_onesheet(t, hints)
    qa = onesheet_qa(os_)
    drawing = draw_onesheet(t, os_, Style(proj.p("rendering")), content)
    pdf_path = proj.out / "pfd_one_sheet_A0.pdf"
    save_onesheet_pdf(drawing, t, pdf_path)
    xml_path = proj.out / "pfd.proteus.xml"
    export_topology(t, xml_path)
    gstats = add_graphics(xml_path, drawing, os_, t, None)
    rt = verify_roundtrip(t, xml_path)
    g_ok = all(ok for _, ok, _ in verify_graphics(xml_path, gstats))
    wb = Workbook()
    ws = wb.active
    ws.title = "Stream table"
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2B2B2B")
    ws2 = wb.create_sheet("Equipment summary")
    ws2.append(["Tag", "Description", "Key data (as stated)"])
    for r in pfd.equipment_summary:
        ws2.append(list(r))
    xlsx = proj.out / "pfd_stream_table.xlsx"
    wb.save(xlsx)
    md = ["# Process flow diagram", "", f"Derived from `source_topology.json` (no P&ID rules). Streams: {len(pfd.streams)}; "
          f"equipment symbols: {len(t.equipment)}; parallel banks drawn once: {len(pfd.banks)}; geometric problems: {len(qa)}; "
          f"DEXPI round trip: {'PASS' if rt.ok and g_ok else 'FAIL'}.", "",
          "Left out (P&ID detail): " + ", ".join(f"{n} {k}" for k, n in sorted(pfd.dropped.items()) if n) + ".", "",
          "## Banks", ""] + [f"- {b.rep}: {', '.join(b.members)}" for b in pfd.banks] + ["", "## Streams", "",
          "| No. | Service | From | To | Temperature | Pressure | Flow |", "|---|---|---|---|---|---|---|"]
    for s_ in pfd.streams:
        v = s_.values
        md.append(f"| {s_.no} | {v.get('Service', '')} | {v.get('From', '')} | {v.get('To', '')} | {v.get('Temperature', '')} | "
                  f"{v.get('Pressure', '')} | {v.get('Flow', '')}{'*' if 'Flow' in s_.derived else ''} |")
    md += ["", "## Layout", ""] + ([f"- {q}" for q in qa] or ["No geometric problems."])
    (proj.reports / "pfd.md").write_text("\n".join(md) + "\n", encoding="utf-8")
    return PFDResult(pfd, qa, drawing.scale, rt.ok and g_ok, [pdf_path, xml_path, xlsx, proj.reports / "pfd.md"])


class SourceChanged(RuntimeError):
    pass


@dataclass
class Project:
    root: Path
    cfg: dict[str, Any]

    def p(self, key: str) -> Path:
        return (self.root / self.cfg[key]).resolve()

    @property
    def out(self) -> Path:
        d = self.p("output_dir")
        d.mkdir(parents=True, exist_ok=True)
        return d

    @property
    def reports(self) -> Path:
        d = self.p("reports_dir")
        d.mkdir(parents=True, exist_ok=True)
        return d

    @property
    def pdf_path(self) -> Path:
        return self.out / "pid_review_draft.pdf"

    @property
    def xml_path(self) -> Path:
        return self.out / "pid_draft.proteus.xml"

    @property
    def sheet_path(self) -> Path:
        return self.out / "pid_one_sheet_A0.pdf"

    @property
    def svg_path(self) -> Path:
        return self.out / "pid_one_sheet_from_dexpi.svg"

    @property
    def lists_path(self) -> Path:
        return self.out / "engineering_lists.xlsx"

    @property
    def plan_path(self) -> Path:
        return self.root / self.cfg.get("review_plan", "review_plan.yaml")


def load_project(path: str | Path) -> Project:
    path = Path(path)
    if path.is_dir():
        path = path / "project.yaml"
    elif not path.exists() and (Path("projects") / str(path) / "project.yaml").exists():
        path = Path("projects") / str(path) / "project.yaml"
    cfg = yaml.safe_load(path.read_text(encoding="utf-8"))
    return Project(root=path.parent.resolve(), cfg=cfg)


def load_policy_for(proj: Project) -> Policy:
    return load_policy(proj.p("generic_policy"), proj.p("rules"))


def _hints(proj: Project) -> dict[str, Any]:
    hp = proj.p("layout_hints")
    return yaml.safe_load(hp.read_text(encoding="utf-8")) if hp.exists() else {}


# ------------------------------------------------------------------------------ stages
def load_working(proj: Project, rebuild: bool = False) -> Topology:
    """Working topology, created from the source snapshot on the first run."""
    src_path, work_path = proj.p("source_topology"), proj.p("working_topology")
    sha = sha256_file(src_path)
    if work_path.exists() and not rebuild:
        topo = load_topology(work_path)
        if topo.metadata.source_snapshot_sha256 != sha:
            raise SourceChanged(
                "source_topology.json changed since working_topology.json was created "
                f"({(topo.metadata.source_snapshot_sha256 or '')[:10]}… -> {sha[:10]}…). "
                "Re-run with --rebuild to start a new working topology (approvals are lost).")
        return topo
    topo = load_topology(src_path)
    topo.metadata.source_snapshot_sha256 = sha
    return topo


def stage_enrich(proj: Project, topo: Topology, policy: Policy) -> EnrichResult:
    res = enrich(topo, policy)
    save_topology(topo, proj.p("working_topology"))
    return res


def check_idempotent(topo: Topology, policy: Policy) -> bool:
    twin = copy.deepcopy(topo)
    enrich(twin, policy)
    return dumps_topology(twin) == dumps_topology(topo)


@dataclass
class RunResult:
    topology: Topology
    enrichment: EnrichResult
    validation: ValidationReport
    export: ExportResult
    roundtrip: RoundTripReport
    layout: DrawingLayout
    qa: dict[str, list[str]]
    scales: dict[str, float]
    idempotent: bool
    onesheet: OneSheet | None = None
    sheet_qa: list[str] | None = None
    register: Register | None = None
    sheet_scale: float = 0.0
    graphics: dict[str, int] | None = None
    extraction_check: list[tuple[str, str, str]] | None = None
    pfd: PFDResult | None = None


def run_all(project: str | Path, rebuild: bool = False, render: bool = True) -> RunResult:
    proj = load_project(project)
    policy = load_policy_for(proj)
    topo = load_working(proj, rebuild)
    enr = stage_enrich(proj, topo, policy)
    idem = check_idempotent(topo, policy)
    val = validate(topo)
    (proj.reports / "validation.md").write_text(validation_md(val), encoding="utf-8")
    if not val.ok:
        raise RuntimeError(f"validation failed with {len(val.errors)} error(s); see {proj.reports / 'validation.md'}")
    exp = export_topology(topo, proj.xml_path)
    hints = _hints(proj)
    # the one-sheet drawing: laid out, recorded once, written to the DEXPI file and to the PDF
    reg = build_register(topo, load_plan(proj.plan_path))
    os_ = build_onesheet(topo, hints)
    sheet_qa = onesheet_qa(os_)
    drawing = draw_onesheet(topo, os_, Style(proj.p("rendering")), sheet_content(topo, reg))
    gstats = add_graphics(proj.xml_path, drawing, os_, topo, None)
    rt = verify_roundtrip(topo, proj.xml_path)
    for name, ok, detail in verify_graphics(proj.xml_path, gstats):
        rt.add(name, ok, detail)
    (proj.reports / "roundtrip.md").write_text(roundtrip_md(rt, proj.xml_path.name), encoding="utf-8")
    dl = build_layouts(topo, hints)
    qa = {s.sheet_id: layout_qa(s) for s in dl.sheets}
    scales: dict[str, float] = {}
    if render:
        save_onesheet_pdf(drawing, topo, proj.sheet_path)
        render_svg(proj.xml_path, proj.svg_path)
        facts = {
            "validation_ok": val.ok, "validation_errors": len(val.errors), "validation_warnings": len(val.warnings),
            "roundtrip_ok": rt.ok, "roundtrip_summary": f"{len(rt.checks)} checks over {rt.stats.get('lines', '?')} lines in {rt.stats.get('segments', '?')} segments",
            "idempotent": "verified in this run" if idem else "NOT idempotent",
        }
        scales = render_review_pdf(topo, dl, proj.pdf_path, proj.p("rendering"), facts)
    write_xlsx(topo, reg, proj.lists_path)
    from .scaffold import check_project

    ext = check_project(proj.root)
    (proj.reports / "extraction_check.md").write_text(
        "# Extraction check\n\nVerbatim quotes, structure and open-point triage of the extraction (`pid check`).\n\n"
        + ("\n".join(f"- **{l}** {w} — fix: {f}" for l, w, f in ext) if ext else "No problems.") + "\n", encoding="utf-8")
    (proj.reports / "decision_register.md").write_text(register_md(reg, topo), encoding="utf-8")
    # reports
    rep = proj.reports
    (rep / "enrichment.md").write_text(R.enrichment_md(enr.log), encoding="utf-8")
    (rep / "proposals.md").write_text(R.proposals_md(topo), encoding="utf-8")
    (rep / "unresolved_requirements.md").write_text(R.unresolved_md(topo), encoding="utf-8")
    (rep / "assumptions.md").write_text(R.assumptions_md(topo, policy), encoding="utf-8")
    (rep / "operating_modes.md").write_text(R.modes_md(topo, policy), encoding="utf-8")
    (rep / "export_limitations.md").write_text(R.export_md(exp.report, proj.xml_path.name), encoding="utf-8")
    (rep / "layout.md").write_text(R.layout_md(qa, scales, sheet_qa, drawing.scale), encoding="utf-8")
    facts_tbl = {
        "Equipment": len(topo.equipment), "Process lines": len(topo.process_lines), "Manifolds (tees)": sum(1 for c in topo.components if c.kind == "manifold"),
        "Valves from the source / proposed": f"{sum(1 for c in topo.components if c.kind != 'manifold' and c.provenance.origin == 'source')} / "
                                           f"{sum(1 for c in topo.components if c.kind != 'manifold' and c.provenance.origin == 'generated')}",
        "Instruments (proposed)": len(topo.instruments), "Control loops": len(topo.control_loops),
        "Pending review": sum(1 for o in topo.objects().values() if o.provenance.origin == "generated" and o.provenance.status == "pending_review"),
        "Open points": sum(1 for u in topo.unresolved_requirements if u.status == "open"),
        "Decisions needed": f"{len(reg.holds)} holds + {len(reg.rules)} rule approvals ({len(reg.datasheet)} datasheet items do not change the drawing)",
        "Enrichment idempotent": "yes" if idem else "NO",
    }
    (rep / "summary.md").write_text(R.summary_md(topo, val, rt, qa, facts_tbl), encoding="utf-8")
    pfd_res = run_pfd(proj.root) if render else None
    return RunResult(topo, enr, val, exp, rt, dl, qa, scales, idem, os_, sheet_qa, reg, drawing.scale, gstats, ext, pfd_res)
