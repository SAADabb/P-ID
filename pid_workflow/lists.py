"""Engineering lists derived from the working topology: equipment list, line list, valve list,
instrument index, holds and datasheet data - one workbook, one sheet per list."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .decisions import Register
from .schema import Topology, is_pending, logical

HEAD = PatternFill("solid", fgColor="2B2B2B")
PROPOSED = PatternFill("solid", fgColor="FFF0D9")


def _status(o: Any) -> str:
    pv = o.provenance
    if pv.origin == "source":
        return "specification"
    return {"pending_review": "proposed – pending review", "approved": "approved"}.get(pv.status or "", pv.status or "")


def _attrs(o: Any) -> str:
    out = []
    for k, a in sorted(o.attributes.items()):
        v = a.value
        if isinstance(v, dict):
            v = "–".join(str(x) for x in v.values())
        out.append(f"{k.replace('_', ' ')}: {v}{(' ' + a.unit) if a.unit else ''}")
    return "; ".join(out)


def _src(o: Any) -> str:
    pv = o.provenance
    if pv.origin == "source":
        return f"p.{pv.page} §{pv.section}"
    return f"rule {pv.rule}"


def build_lists(topo: Topology, reg: Register) -> dict[str, tuple[list[str], list[list[Any]], list[bool]]]:
    objs = topo.objects()

    def name_of(ref: str) -> str:
        o = objs.get(ref)
        if o is None:
            return ref
        tag = getattr(o, "tag", None)
        return f"{tag} {o.name or ''}".strip() if tag else (o.name or ref)

    sheets: dict[str, tuple[list[str], list[list[Any]], list[bool]]] = {}
    rows, prop = [], []
    for e in sorted(topo.equipment, key=lambda x: (x.tag or "~", x.id)):
        rows.append([e.tag or "TBD", e.tag_status, e.name, e.model_class, e.area, e.lifecycle, _attrs(e), len(e.nozzles), _src(e), e.id])
        prop.append(False)
    sheets["Equipment list"] = (["Tag", "Tag status", "Description", "DEXPI class", "Area", "Lifecycle", "Data from the specification", "Nozzles", "Source", "Topology id"], rows, prop)

    groups: dict[str, list[Any]] = {}
    for ln in topo.process_lines:
        groups.setdefault(logical(ln), []).append(ln)
    rows, prop = [], []
    for lg, lns in sorted(groups.items(), key=lambda kv: (kv[1][0].line_number or "~", kv[0])):
        tos = {x.to.ref for x in lns}
        froms = {x.from_.ref for x in lns}
        first = next((x for x in lns if x.from_.ref not in tos), lns[0])
        last = next((x for x in lns if x.to.ref not in froms), lns[-1])
        p = first.properties
        size = ""
        if p.get("nominal_diameter") and p["nominal_diameter"].value:
            size = f"DN{p['nominal_diameter'].value}"
        elif p.get("nominal_diameter_range"):
            size = f"DN{p['nominal_diameter_range'].value['min']}–{p['nominal_diameter_range'].value['max']}"
        elif p.get("nominal_diameter_min"):
            size = f"≥DN{p['nominal_diameter_min'].value}"
        others = "; ".join(f"{k.replace('_', ' ')}: {a.value}{(' ' + a.unit) if a.unit else ''}" for k, a in sorted(p.items()) if not k.startswith("nominal_diameter"))
        rows.append([first.line_number or "", first.line_number_status, first.service, name_of(first.from_.ref), name_of(last.to.ref), size or "TBD",
                     "TBD", others, len(lns), _src(first), lg])
        prop.append(first.line_number_status == "proposed")
    sheets["Line list"] = (["Line No.", "No. status", "Service", "From", "To", "Size", "Piping class", "Other data", "Pieces", "Source", "Logical line"], rows, prop)

    line_of: dict[str, str] = {}
    for ln in topo.process_lines:
        for end in (ln.from_.ref, ln.to.ref):
            line_of.setdefault(end, ln.line_number or ln.id)
    rows, prop = [], []
    for c in sorted([c for c in topo.components if c.kind != "manifold"], key=lambda x: (x.tag or "~", x.id)):
        rows.append([c.tag or "", c.function or "", c.model_class, c.technology or "to be selected", c.operation, line_of.get(c.id, ""),
                     c.open_end or "", _status(c), _src(c), c.id])
        prop.append(is_pending(c.provenance))
    sheets["Valve list"] = (["Tag", "Function", "DEXPI class", "Technology", "Operation", "Line", "Open end", "Status", "Source", "Topology id"], rows, prop)

    rows, prop = [], []
    for i in sorted(topo.instruments, key=lambda x: (x.tag or "~", x.id)):
        t = i.target or {}
        where = name_of(t.get("ref", "")) if t.get("kind") == "equipment" else (next((ln.line_number for ln in topo.process_lines if ln.id == t.get("ref")), None) or t.get("ref", ""))
        rows.append([i.tag or "", i.isa_code, i.measured_variable, i.function, where, i.loop_id or "", "TBD", "TBD", _status(i), _src(i), i.id])
        prop.append(is_pending(i.provenance))
    sheets["Instrument index"] = (["Tag", "ISA code", "Variable", "Function", "Location / line", "Loop", "Range", "Alarm / trip", "Status", "Source", "Topology id"], rows, prop)

    rows = [[f"H{r.hold.num}", r.hold.title, r.drawn_as, r.question, ", ".join(r.hold.affects), "", r.hold.ur_id] for r in reg.holds]
    rows += [["closed", r.hold.title, r.drawn_as, r.question, ", ".join(r.hold.affects), r.answer, r.hold.ur_id] for r in reg.closed]
    sheets["Holds"] = (["Hold", "Decision", "Drawn today", "Question", "Affects", "Answer", "Open point"], rows, [False] * len(rows))
    rows = [[rid, n, d, f"pid review <project> --approve-rule {rid}"] for rid, n, d in reg.rules]
    sheets["Rule approvals"] = (["Rule", "Pending objects", "Adds", "Command"], rows, [False] * len(rows))
    rows = [[uid, t] for uid, t in reg.datasheet]
    sheets["Datasheet data"] = (["Open point", "Missing data"], rows, [False] * len(rows))
    return sheets


def write_xlsx(topo: Topology, reg: Register, path: str | Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    for title, (head, rows, prop) in build_lists(topo, reg).items():
        ws = wb.create_sheet(title)
        ws.append(head)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = HEAD
            c.alignment = Alignment(vertical="center", wrap_text=True)
        for r, p in zip(rows, prop):
            ws.append(r)
            if p:
                for c in ws[ws.max_row]:
                    c.fill = PROPOSED
        for i, h in enumerate(head, 1):
            width = max([len(str(h))] + [min(len(str(r[i - 1])), 60) for r in rows]) + 2
            ws.column_dimensions[get_column_letter(i)].width = min(max(width, 8), 62)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(vertical="top", wrap_text=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    wb.properties.creator = "pid_workflow"
    wb.save(path)
    return path
